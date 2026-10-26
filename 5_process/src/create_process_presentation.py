"""Create a concise, evidence-linked Process presentation without article text."""
from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[2]
PROCESS = REPO / "5_process"
RAW = REPO / "1_du_lieu_chinh" / "vn_labor_news_dataset.csv"
REPORTS = PROCESS / "reports"
CHARTS = REPORTS / "charts"
OUTPUT_DEFAULT = REPORTS / "presentation.pptx"
NOTES_DEFAULT = REPORTS / "presentation_notes.md"

NAVY = RGBColor(22, 43, 74)
BLUE = RGBColor(42, 104, 173)
TEAL = RGBColor(0, 133, 137)
GOLD = RGBColor(226, 166, 59)
INK = RGBColor(38, 51, 67)
MUTED = RGBColor(93, 108, 125)
PALE = RGBColor(240, 245, 250)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(171, 58, 58)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=INK, spacing=9):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.space_after = Pt(spacing)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
    return box


def add_card(slide, x, y, w, h, heading, body, accent=BLUE, body_size=16):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PALE
    shape.line.fill.background()
    shape.text_frame.clear()
    shape.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    shape.text_frame.margin_left = Inches(0.18)
    shape.text_frame.margin_right = Inches(0.16)
    shape.text_frame.margin_top = Inches(0.16)
    shape.text_frame.margin_bottom = Inches(0.1)
    p = shape.text_frame.paragraphs[0]
    p.space_after = Pt(8)
    run = p.add_run()
    run.text = heading
    run.font.name = "Aptos Display"
    run.font.size = Pt(19)
    run.font.bold = True
    run.font.color.rgb = accent
    p2 = shape.text_frame.add_paragraph()
    p2.text = body
    p2.font.name = "Aptos"
    p2.font.size = Pt(body_size)
    p2.font.color.rgb = INK
    return shape


def add_base(prs, title, subtitle, source, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(slide, title, 0.65, 0.34, 12.0, 0.52, 26, NAVY, True, "Aptos Display")
    if subtitle:
        add_text(slide, subtitle, 0.68, 0.91, 12.0, 0.38, 12, MUTED)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(1.35), Inches(11.95), Inches(0.035)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()
    add_text(slide, source, 0.68, 7.15, 11.5, 0.2, 8, MUTED)
    add_text(slide, str(page), 12.1, 7.12, 0.45, 0.24, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_chart(slide, filename, x, y, w, h):
    path = CHARTS / filename
    if not path.is_file():
        raise FileNotFoundError(f"Required generated chart is missing: {path}")
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def build_presentation(output: Path, notes_path: Path) -> None:
    raw = pd.read_csv(
        RAW,
        usecols=["article_id", "publish_date", "source_name", "topic_primary", "sentiment"],
        encoding="utf-8-sig",
        dtype={"article_id": str},
    )
    profile = json.loads((REPORTS / "00_input_profile.json").read_text(encoding="utf-8"))
    integration = json.loads((REPORTS / "integration_validation.json").read_text(encoding="utf-8"))
    split_report = json.loads((REPORTS / "split_and_leakage_checks.json").read_text(encoding="utf-8"))
    sample = pd.read_csv(
        PROCESS / "data" / "processed" / "manual_label_sample.csv",
        usecols=[
            "relevance_manual",
            "topic_manual",
            "review_reason",
            "reviewer_id",
            "review_date",
        ],
        keep_default_na=False,
        encoding="utf-8-sig",
    )
    manual_fields = list(sample.columns)
    complete_review = sample[manual_fields].apply(
        lambda column: column.astype(str).str.strip().ne("")
    ).all(axis=1)
    manual_reviewed = int(complete_review.sum())
    sample_size = len(sample)
    manual_pending = sample_size - manual_reviewed
    train = split_report["split_counts"]["train"]
    validation = split_report["split_counts"]["validation"]
    test = split_report["split_counts"]["test"]
    rows = len(raw)
    if rows != profile["rows"] or integration["status"] != "PASS":
        raise ValueError("Presentation inputs are not aligned with the validated integration report.")
    if train + validation + test != rows:
        raise ValueError("Split totals do not equal the current source row count.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "PROCESS", 0.85, 0.75, 5.0, 0.55, 18, RGBColor(115, 211, 207), True)
    add_text(slide, "VN Labor News Dataset", 0.85, 1.65, 11.6, 0.9, 34, WHITE, True, "Aptos Display")
    add_text(slide, "Từ dữ liệu thu thập đến bộ dữ liệu có kiểm tra và truy vết", 0.9, 2.75, 10.9, 0.65, 21, WHITE)
    add_text(slide, f"{rows:,} bài  |  8 nguồn báo  |  01/2022–10/2026", 0.9, 4.0, 10.0, 0.5, 19, RGBColor(220, 231, 244), True)
    add_text(slide, "Báo cáo tích hợp Process • 27/10/2026", 0.9, 6.65, 10.5, 0.28, 11, RGBColor(190, 205, 221))

    slide = add_base(prs, "Mục tiêu và phạm vi", "Process chuẩn hóa và kiểm định dữ liệu — chưa huấn luyện mô hình.", "Nguồn: Phân công Process 6 người; 00_input_profile.json; process_quality_report.md", 2)
    add_card(slide, 0.75, 1.7, 3.75, 2.0, "Đầu vào chuẩn", "CSV gốc 12.784 dòng, 16 cột; 8 nguồn báo; ngày đăng từ 2022-01-01 đến 2026-10-23.", BLUE)
    add_card(slide, 4.8, 1.7, 3.75, 2.0, "Nguyên tắc dữ liệu", "Raw chỉ đọc. Chuẩn hóa và làm sạch ở đầu ra riêng; giữ ID để truy vết; không xóa bài nghi trùng.", TEAL)
    add_card(slide, 8.85, 1.7, 3.75, 2.0, "Giới hạn", "Nhãn chủ đề tự động là weak labels; sentiment đang trống; địa danh là thông tin được trích, không khẳng định nơi sự kiện xảy ra.", GOLD, 15)
    add_text(slide, "Mục tiêu bàn giao: pipeline chạy lại được, kết quả có bằng chứng kiểm tra, rủi ro chưa giải quyết được nêu rõ.", 1.0, 4.45, 11.1, 0.8, 22, NAVY, True, align=PP_ALIGN.CENTER)

    slide = add_base(prs, "Pipeline và các cổng kiểm tra", "Mọi đầu ra đều liên kết lại bằng article_id; CSV nguồn không bị ghi đè.", "Nguồn: README_PROCESS.md; integration_validation.json", 3)
    stages = [
        ("01", "Profile + schema", "12.784 × 16\nSchema PASS", BLUE),
        ("02", "Metadata", "Chuẩn hóa + audit\nkhông xóa bài", TEAL),
        ("03", "Text cleaning", "title_clean\ncontent_clean", GOLD),
        ("04", "Thống kê + split", "3 manifest theo thời gian\nleakage audit", BLUE),
        ("05", "Tích hợp", "Đối chiếu raw, ID,\ntext, split", TEAL),
    ]
    x0, gap, width = 0.72, 0.18, 2.28
    for index, (number, heading, body, accent) in enumerate(stages):
        x = x0 + index * (width + gap)
        add_card(slide, x, 2.0, width, 2.25, f"{number}  {heading}", body, accent, 15)
        if index < len(stages) - 1:
            add_text(slide, "→", x + width + 0.005, 2.82, 0.18, 0.35, 21, MUTED, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Cổng cuối: 45/45 test Process + 7/7 test gốc; integration validation = PASS.", 1.0, 5.05, 11.0, 0.55, 21, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(slide, "Kết quả tích hợp xác nhận 12.784 ID metadata, 12.784 dòng text-clean và 12.784 ID split khớp nguồn.", 1.0, 5.72, 11.0, 0.55, 16, INK, align=PP_ALIGN.CENTER)

    slide = add_base(prs, "Hồ sơ dữ liệu và kiểm tra chất lượng", "Raw preservation: checksum trước/sau giống nhau.", "Nguồn: 00_input_profile.json; 01_schema_validation.json; integration_validation.json", 4)
    add_card(slide, 0.75, 1.7, 3.45, 1.55, "SHA-256", profile["sha256"][:16] + "…", NAVY, 19)
    add_card(slide, 4.45, 1.7, 3.45, 1.55, "Ngày xuất bản", "2022-01-01 → 2026-10-23", BLUE, 18)
    add_card(slide, 8.15, 1.7, 4.15, 1.55, "Schema", "PASS • 0 lỗi Critical/High", TEAL, 18)
    add_chart(slide, "articles_by_source.png", 0.65, 3.35, 6.0, 3.35)
    add_bullets(slide, [
        "8 nguồn nằm trong danh sách chuẩn; article_id và source_url duy nhất.",
        "topic_secondary trống 5.746 dòng; sentiment trống 12.784 dòng — giữ nguyên trạng thái thiếu.",
        "word_count khác content.split() ở 11.953 dòng do khác định nghĩa đầu vào; không ghi đè giá trị raw.",
    ], 6.95, 3.7, 5.45, 2.7, 15, INK, 11)

    slide = add_base(prs, "Metadata và kiểm định nhãn", "Tách rõ kiểm tra tự động, đề xuất rà soát và xác nhận của con người.", "Nguồn: metadata_audit.md; label_quality.md; issue_log.csv", 5)
    add_card(slide, 0.8, 1.75, 3.7, 1.6, "Metadata", "14 thay đổi chuẩn hóa được ghi log; 9 lượt ngoại lai được đánh cờ; không xóa bản ghi.", BLUE, 15)
    add_card(slide, 4.8, 1.75, 3.7, 1.6, "Audit trùng", "71 dòng review; 62 cụm metadata. Đây là ứng viên audit, không phải phán quyết xóa.", TEAL, 15)
    add_card(slide, 8.8, 1.75, 3.7, 1.6, "Nhãn thủ công", f"{manual_reviewed}/{sample_size} đã được con người xác nhận.\nCòn {manual_pending} dòng chờ review.", GOLD, 15)
    add_chart(slide, "articles_by_topic.png", 0.85, 3.45, 7.15, 3.4)
    add_bullets(slide, [
        "Phân bố A–L mô tả nhãn tự động hiện có, không phải chất lượng nhãn.",
        "Không báo cáo precision/recall hay độ đồng thuận khi chưa có nhãn người.",
        "Các gợi ý AI không được dùng làm kết quả review chính thức.",
    ], 8.35, 3.8, 4.0, 2.6, 15, INK, 12)

    slide = add_base(prs, "Thống kê mô tả", "Độ dài được tính từ content theo whitespace split; không phải tokenizer ngôn ngữ.", "Nguồn: process_statistics.md; charts được sinh từ create_splits_and_leakage_audit.py", 6)
    add_chart(slide, "articles_by_year_month.png", 0.55, 1.65, 7.4, 4.95)
    add_chart(slide, "content_word_count_distribution.png", 8.0, 1.8, 4.8, 3.3)
    add_card(slide, 8.3, 5.25, 4.1, 1.15, "Cách diễn giải", "Số lượng theo thời gian và độ dài để mô tả dữ liệu; không suy ra tính đại diện xã hội.", TEAL, 13)

    slide = add_base(prs, "Split theo thời gian và kiểm tra leakage", "Split là manifest mặc định để bàn giao, chưa phải quyết định cuối cho đánh giá mô hình.", "Nguồn: split_validation.md; split_and_leakage_checks.json; handoff_to_modeling.md", 7)
    split_cards = [("TRAIN", train, "Trước 2025", BLUE), ("VALIDATION", validation, "Năm 2025", TEAL), ("TEST", test, "Từ 2026", GOLD)]
    for i, (name, count, period, accent) in enumerate(split_cards):
        add_card(slide, 0.8 + i * 4.1, 1.75, 3.65, 1.55, name, f"{count:,} bài • {period}", accent, 18)
    add_card(slide, 0.8, 3.75, 5.7, 1.8, "Ứng viên tiêu đề", "118 cặp được sàng lọc; 32 cặp nằm khác split. Không đồng nghĩa đã xác nhận trùng.", BLUE, 16)
    add_card(slide, 6.8, 3.75, 5.7, 1.8, "Cụm metadata", "2 cụm giao split: DUP-0037 (train/validation), DUP-0059 (validation/test). 65 ID liên quan được giữ ở review_hold.", RED, 15)
    add_text(slide, "Chỉ dùng các dòng eligible; review và ghi quyết định trước khi đưa holdout trở lại.", 1.0, 6.05, 11.0, 0.45, 17, NAVY, True, align=PP_ALIGN.CENTER)

    slide = add_base(prs, "Trạng thái bàn giao theo phân công", "Phân biệt hoàn tất kỹ thuật với phần cần xác nhận hoặc bàn giao từ thành viên.", "Nguồn: integration_checklist.md; process_quality_report.md; issue_log.csv", 8)
    status_rows = [
        ("Người 1", "Profile, schema, tích hợp và đối chiếu", "Hoàn tất kỹ thuật", TEAL),
        ("Người 2", "Metadata cleaning + danh sách audit", "Tích hợp; không xóa bài", BLUE),
        ("Người 3", "Text cleaning + tests + output", "12.784 dòng đã kiểm tra", TEAL),
        ("Người 4", "Hướng dẫn + mẫu + validator", f"{manual_pending} nhãn chờ người xác nhận", GOLD),
        ("Người 5", "Thống kê + split + leakage audit", "65 ID đã biết ở review_hold", GOLD),
        ("Người 6", "Slide + ghi chú + nguồn số liệu", "Đang bàn giao trong bộ này", TEAL),
    ]
    y = 1.65
    for who, deliverable, status, accent in status_rows:
        add_text(slide, who, 0.9, y, 1.5, 0.35, 16, accent, True)
        add_text(slide, deliverable, 2.4, y, 5.6, 0.35, 15, INK)
        add_text(slide, status, 8.2, y, 4.2, 0.35, 15, accent, True)
        y += 0.72
    add_text(slide, "Điều kiện còn lại: người review xác nhận nhãn; giữ holdout ngoài modeling đến khi có quyết định.", 1.0, 6.25, 11.1, 0.48, 16, NAVY, True, align=PP_ALIGN.CENTER)

    slide = add_base(prs, "Kết luận và giới hạn", "Kết quả Process có thể tái tạo; chưa đồng nghĩa đã có bộ nhãn vàng hay mô hình được đánh giá.", "Nguồn: process_quality_report.md; labeling_guideline.md; handoff_to_modeling.md", 9)
    add_bullets(slide, [
        "Raw được bảo toàn; các bước xử lý có đầu ra riêng, kiểm thử và kiểm tra tích hợp.",
        f"Nhãn tự động A–L là weak labels; {manual_pending} mẫu chờ xác nhận người thật. Không dùng gợi ý AI như ground truth.",
        "Sentiment chưa được gán; địa danh là nơi được nhắc trong nội dung, không mặc định là địa điểm sự kiện.",
        "Dữ liệu báo chí không đại diện cho toàn xã hội; phân bố nguồn/chủ đề có thể lệch theo thời gian.",
        "Không huấn luyện mô hình trong giai đoạn Process; cần rà soát leakage trước khi chốt đánh giá cuối.",
    ], 1.05, 1.8, 11.2, 3.8, 19, INK, 15)
    add_text(slide, "Không có toàn văn bài báo trong slide.", 1.0, 6.2, 11.1, 0.45, 18, TEAL, True, align=PP_ALIGN.CENTER)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    notes = [
        "# Ghi chú thuyết trình — Process VN Labor News Dataset",
        "",
        "Slide được sinh bằng `python 5_process/src/create_process_presentation.py` từ báo cáo Process và CSV nguồn ở chế độ chỉ đọc. Không chèn tiêu đề hoặc toàn văn bài báo.",
        "",
        "## Nguồn số liệu theo slide",
        "",
        "1. Tiêu đề: `Phan_cong_Process_6_nguoi.md`, `00_input_profile.json`.",
        "2. Mục tiêu/phạm vi: `Phan_cong_Process_6_nguoi.md`, `00_input_profile.json`, `process_quality_report.md`.",
        "3. Pipeline: `README_PROCESS.md`, `integration_validation.json`.",
        "4. Chất lượng: `00_input_profile.json`, `01_schema_validation.json`, `integration_validation.json`, `process_statistics.md`.",
        "5. Metadata/nhãn: `metadata_audit.md`, `label_quality.md`, `issue_log.csv`.",
        "6. Thống kê: `process_statistics.md` và biểu đồ do `create_splits_and_leakage_audit.py` sinh.",
        "7. Split/leakage: `split_validation.md`, `split_and_leakage_checks.json`, `handoff_to_modeling.md`.",
        "8. Tiến độ thành viên: `integration_checklist.md`, `process_quality_report.md`, `issue_log.csv`.",
        "9. Kết luận/hạn chế: `process_quality_report.md`, `labeling_guideline.md`, `handoff_to_modeling.md`.",
        "",
        "## Lời dẫn gợi ý",
        "",
        "- Nhấn mạnh checksum và raw preservation: kết quả biến đổi chỉ nằm ở thư mục Process.",
        "- Phân biệt nhãn tự động với nhãn thủ công; mẫu review chưa được người xác nhận.",
        "- Gọi split là phương án mặc định theo thời gian; 65 ID đã biết bị giữ ngoài modeling.",
        "- 32 cặp tiêu đề và 2 cụm duplicate giao split là cảnh báo ứng viên, không khẳng định chắc chắn trùng.",
        "- Không trình bày các bản nháp AI của metadata như quyết định chính thức.",
    ]
    notes_path.parent.mkdir(parents=True, exist_ok=True)
    notes_path.write_text("\n".join(notes) + "\n", encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT_DEFAULT)
    parser.add_argument("--notes", type=Path, default=NOTES_DEFAULT)
    args = parser.parse_args()
    build_presentation(args.output, args.notes)
    print(f"Created presentation: {args.output}")
    print(f"Created speaker notes: {args.notes}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
